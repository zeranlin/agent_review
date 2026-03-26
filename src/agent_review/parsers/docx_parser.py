from __future__ import annotations

from pathlib import Path
import re

from docx import Document

from ..models import ParseResult, ParsedPage, ParsedTable, RawBlock, RawCell, RawTable, SourceAnchor


def parse_docx(path: str | Path) -> ParseResult:
    target = Path(path).expanduser().resolve()
    document = Document(str(target))

    raw_blocks = _parse_docx_blocks(document, target)
    raw_tables, table_records, table_text_blocks = _parse_docx_tables(document, target)
    paragraph_texts = [block.text for block in raw_blocks if block.text.strip()]
    text_parts = paragraph_texts + table_text_blocks
    text = "\n".join(part for part in text_parts if part)
    return ParseResult(
        parser_name="docx",
        source_path=str(target),
        source_format="docx",
        page_count=None,
        text=text,
        pages=[ParsedPage(page_index=1, text=text, source="docx")],
        tables=table_records,
        raw_blocks=raw_blocks,
        raw_tables=raw_tables,
        warnings=[],
    )


def _parse_docx_blocks(document: Document, target: Path) -> list[RawBlock]:
    blocks: list[RawBlock] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = _paragraph_style_name(paragraph)
        numbering = _extract_numbering_text(text)
        metadata = {
            "heading_candidate": _is_heading_candidate(text, style_name),
            "catalog_candidate": _is_catalog_candidate(text),
            "style_name": style_name,
            "numbering_level_guess": _guess_numbering_level(text, style_name),
        }
        blocks.append(
            RawBlock(
                block_id=f"p-{paragraph_index}",
                block_type="paragraph",
                text=text,
                style_name=style_name,
                numbering=numbering,
                anchor=SourceAnchor(
                    source_path=str(target),
                    block_no=paragraph_index,
                    paragraph_no=paragraph_index,
                ),
                metadata=metadata,
            )
        )
    return blocks


def _parse_docx_tables(
    document: Document,
    target: Path,
) -> tuple[list[RawTable], list[ParsedTable], list[str]]:
    raw_tables: list[RawTable] = []
    parsed_tables: list[ParsedTable] = []
    table_text_blocks: list[str] = []
    for table_index, table in enumerate(document.tables, start=1):
        raw_rows: list[list[RawCell]] = []
        parsed_rows: list[list[str]] = []
        for row_index, row in enumerate(table.rows, start=1):
            raw_row: list[RawCell] = []
            parsed_row: list[str] = []
            for col_index, cell in enumerate(row.cells, start=1):
                cell_text = cell.text.strip()
                raw_row.append(
                    RawCell(
                        row_index=row_index,
                        col_index=col_index,
                        text=cell_text,
                        is_header=row_index == 1,
                        anchor=SourceAnchor(
                            source_path=str(target),
                            table_no=table_index,
                            row_no=row_index,
                            cell_no=col_index,
                        ),
                    )
                )
                parsed_row.append(cell_text)
            if any(item for item in parsed_row):
                raw_rows.append(raw_row)
                parsed_rows.append(parsed_row)
        if not parsed_rows:
            continue
        raw_tables.append(
            RawTable(
                table_id=f"t-{table_index}",
                rows=raw_rows,
                anchor=SourceAnchor(source_path=str(target), table_no=table_index),
                title_hint=_infer_table_title(parsed_rows),
                metadata={
                    "header_row_count": 1,
                    "column_count": max((len(row) for row in parsed_rows), default=0),
                },
            )
        )
        parsed_tables.append(
            ParsedTable(
                table_index=table_index,
                row_count=len(parsed_rows),
                rows=parsed_rows,
                source="docx_table",
            )
        )
        table_text_blocks.append("\n".join(" | ".join(cell for cell in row) for row in parsed_rows))
    return raw_tables, parsed_tables, table_text_blocks


def _paragraph_style_name(paragraph) -> str:
    style = getattr(paragraph, "style", None)
    return getattr(style, "name", "") or ""


def _extract_numbering_text(text: str) -> str:
    match = re.match(r"^\s*((?:第[一二三四五六七八九十百千0-9]+[章节册部分篇])|(?:[一二三四五六七八九十]+、)|(?:（[一二三四五六七八九十0-9]+）)|(?:\d+[\.、]))", text)
    return match.group(1) if match else ""


def _is_heading_candidate(text: str, style_name: str) -> bool:
    if style_name and "heading" in style_name.lower():
        return True
    return bool(_extract_numbering_text(text)) or bool(
        re.match(r"^(目录|关键信息|项目概况|投标人资格要求|评分标准|合同条款)", text)
    )


def _is_catalog_candidate(text: str) -> bool:
    if text == "目录":
        return True
    return bool(
        re.match(
            r"^(第[一二三四五六七八九十百千0-9]+[章节册部分篇]|[一二三四五六七八九十]+、|（[一二三四五六七八九十0-9]+）)",
            text,
        )
        and len(text) <= 40
    )


def _guess_numbering_level(text: str, style_name: str) -> int:
    if style_name and "heading 1" in style_name.lower():
        return 1
    if style_name and "heading 2" in style_name.lower():
        return 2
    if re.match(r"^第[一二三四五六七八九十百千0-9]+[章节册部分篇]", text):
        return 1
    if re.match(r"^[一二三四五六七八九十]+、", text):
        return 2
    if re.match(r"^（[一二三四五六七八九十0-9]+）", text):
        return 3
    if re.match(r"^\d+[\.、]", text):
        return 4
    return 0


def _infer_table_title(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    first_row = " ".join(cell for cell in rows[0] if cell).strip()
    return first_row[:80]
